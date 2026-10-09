"""Builds build/backend-doc/TML_Masters_Backend_Design.docx for the Solution Architect (test data only).
Run `npm run backend-doc` (draws the diagrams first). Needs: pip install python-docx."""
import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(os.path.join(HERE, '../../build/backend-doc'))
NAVY = RGBColor(0x00, 0x22, 0x44)

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(0.8)
sec.top_margin = sec.bottom_margin = Inches(0.8)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(10.5)
for name, size in (('Heading 1', 17), ('Heading 2', 13)):
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = NAVY


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)


def para(text='', bold=False, italic=False, size=None, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullets(items):
    for t in items:
        p = doc.add_paragraph(style='List Bullet')
        if isinstance(t, tuple):
            r = p.add_run(t[0])
            r.bold = True
            p.add_run(t[1])
        else:
            p.add_run(t)
        p.paragraph_format.space_after = Pt(2)


def table(head, rows, widths=None, font=9.5):
    t = doc.add_table(rows=1, cols=len(head))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font)
        r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, '002244')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(font)
            if i == 0:
                r.bold = True
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def picture(name, width=6.6):
    doc.add_picture(os.path.join(OUT, name), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


# footer with page number
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run('TML Service Transformation · Masters backend design · draft for review · test data only · page ')
fr.font.size = Pt(8)
for tag, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
    run = fp.add_run()
    run.font.size = Pt(8)
    if tag:
        el = OxmlElement('w:fldChar')
        el.set(qn('w:fldCharType'), tag)
    else:
        el = OxmlElement('w:instrText')
        el.set(qn('xml:space'), 'preserve')
        el.text = txt
    run._r.append(el)

# ------------------------------------------------------------------ title
para('TML Service Transformation', bold=True, size=12, color=RGBColor(0x0B, 0x4F, 0x8A), space_after=2)
para('Masters backend design', bold=True, size=26, color=NAVY, space_after=2)
para('Masters from CRM, masters kept in the ST Portal, and the business control layer', size=14, color=RGBColor(0x47, 0x55, 0x69), space_after=10)
para('For: Solution Architect · Draft for review · 9 October 2026 · Test data only', italic=True, size=10, space_after=12)

doc.add_heading('1. In one page', level=1)
para('The mail of 8 October 2026 sets the approach for master data between CRM and the ST Portal. This document shows how the database supports it, so that it can be reviewed before anything else is built. The design is generic: it does not depend on which masters the BA teams finally send, because a master is data (a definition plus rows), not a table.')
bullets([
    ('Masters already in CRM (for example PL, PPL): ', 'CRM stays the System of Record. The ST Portal reads them through Solar or an API and nobody edits them in ST.'),
    ('New ST masters: ', 'kept in the ST Portal with administration and governance controls. Their lists of values live in the ST Common LOV Master.'),
    ('Business control layer: ', 'Business can switch records on or off, or configure them, for a given ST function, without touching the CRM data. One generic control table serves every master, so a master that Business asks about later needs a setting, not a new table.'),
    ('The one rule that makes this safe: ', 'CRM data and business choices are stored in different places. A CRM refresh can therefore never overwrite a business choice, and a business choice can never change CRM data.'),
])
para('The design is delivered as a database script for PostgreSQL 16 (schema_masters_crm.sql, on top of the existing schema.sql) and is proven by a self-checking test (45 checks and 19 refused wrong actions) that runs on every change. Section 10 lists the decisions we need from the architect.', space_after=8)

doc.add_heading('2. The three kinds of master', level=1)
picture('d1.png')
table(['Kind', 'System of Record', 'In the ST Portal', 'How it is set up in the database'], [
    ['A  CRM master, used as it is', 'CRM', 'Read-only copy, filled by the sync', "source_system = CRM, sync_mode = API or SOLAR, control_mode = NONE"],
    ['B  New ST master', 'ST Portal', 'Created and maintained in ST', 'source_system = ST (as today; nothing changes)'],
    ['C  CRM master + business control', 'CRM', 'Read-only copy + business choices', 'source_system = CRM, control_mode = ENABLE_DISABLE or CONFIGURE, plus one policy row per ST function'],
    ['Common LOV Master', 'Per list', 'Some lists come from CRM, some are ST lists', 'source_system = MIXED; every list type is registered in lov_type with its source'],
], widths=[1.7, 1.1, 1.9, 2.3])

doc.add_heading('3. One record, three separate layers', level=1)
picture('d2.png')
bullets([
    ('CRM part (crm_data): ', 'filled only by the CRM sync. The database refuses any other writer. The record keeps its CRM key; for CRM records the record_id is the CRM key, so references stay stable.'),
    ('ST part (data): ', 'the fields that exist only in ST (for example a remark, or the new fields of an ST module). Users edit only these.'),
    ('Business control (master_control): ', 'one row per record and per ST function: ENABLED, DISABLED or PENDING_REVIEW, optional settings, who decided and why.'),
    ('Field ownership: ', 'every field of a master is flagged CRM or ST. The database refuses a user edit of a CRM field and refuses a CRM feed that tries to fill an ST field. A master can therefore mix CRM fields and new ST fields without conflict.'),
])

doc.add_heading('4. How the data moves', level=1)
picture('d3.png')

doc.add_heading('5. Business control', level=1)
para('A control scope is what Business controls for: normally an ST function (for example "eQC Guided Check"), optionally a module or a dealer. Scopes are rows in control_scope, so a new one needs no deployment. For each master and scope there is one policy row.')
table(['Setting', 'Meaning', 'Example'], [
    ['default_state = ENABLED (deny-list)', 'Everything is available; Business switches some records off.', 'Job card creation offers every PPL model except the ones switched off.'],
    ['default_state = DISABLED (allow-list)', 'Nothing is available until Business enables it. A record that arrives from CRM is stored as PENDING_REVIEW, so Business sees what is new.', 'eQC Guided Check offers only the models Business has enabled.'],
    ['requires_approval + approver_role', 'A change is only requested until a second person approves it. The requester can never approve.', 'Hiding a dealer from dashboards needs approval by the TML Admin.'],
    ['control_mode = CONFIGURE', 'Besides on / off, Business can store settings per record and scope (a JSON object).', 'A daily job-card limit per outlet in one function.'],
    ['control_mode = NONE (default)', 'No control for this master. Switching it on later is an update of this setting plus a policy row.', 'A master Business asks about next month.'],
], widths=[2.0, 2.9, 2.1])
para('What the portal and the dealer apps read: the view master_available (records that are ENABLED for the ST function, with the CRM part and the ST part merged). A record removed in CRM, or made inactive, is never available. master_new_for_review lists the new records waiting for a decision. Every control change is kept in master_control_history (append-only: who, when, why, old and new state) and raises the master version so dealer apps re-synchronise.')

doc.add_heading('6. The CRM sync', level=1)
para('A sync is two steps in the database. The integration job (Solar or API, to be built with the architect) opens a run with start_master_sync and writes the CRM rows into master_sync_stage. Then apply_master_sync applies the run in one transaction. The mapping from CRM field names to master fields is data (master_sync_config.field_mapping).')
table(['What CRM sends', 'What the database does', 'Business choice'], [
    ['A new key', 'Creates the record (record_id = CRM key). In allow-list scopes it waits as PENDING_REVIEW.', 'Decided by Business'],
    ['A changed record', 'Replaces only the CRM part and records old and new values in the history.', 'Kept as it is'],
    ['An unchanged record', 'Writes nothing: no new version, no history entry (compared with a hash).', 'Kept'],
    ['Missing from a FULL run', 'Marks the record REMOVED_IN_CRM. It is never offered to users, but it is kept with its history and its business choice. Nothing is deleted.', 'Kept'],
    ['A removed record comes back', 'Marked ACTIVE again and offered again exactly as before.', 'Back as Business chose'],
    ['A row that cannot be applied', 'Listed per record in master_sync_item; the run ends as PARTIAL and the other rows are applied.', '-'],
    ['DELTA run', 'Applies only what is sent; never marks missing records as removed.', 'Kept'],
], widths=[1.8, 3.7, 1.5])
para('Every run is recorded (master_sync_run): mode, status, counts of new / changed / unchanged / removed / restored / errors, who started it, and the CRM watermark for delta runs.')

doc.add_heading('7. Lists of values and references', level=1)
bullets([
    ('Where the lists live: ', 'the ST Common LOV Master (type = MODULE_FIELD, one row per value). Each list type is registered in lov_type with its source, CRM or ST.'),
    ('A CRM list ', 'is read-only in ST. ST may add its own extra values only when that list has allow_st_extension on; the extra values stay flagged as ST values and never change the CRM ones.'),
    ('A new ST field that uses an existing CRM list: ', 'the field points to the list type (master_field.lov_type) and stores the stable code of the value, not its display text. A list can be used by any number of fields in any master.'),
    ('A field that points to another master ', '(master_field.ref_master) stores the record_id of that master, which for CRM masters is the CRM key.'),
    ('If CRM removes a value that ST records still use: ', 'those records are flagged needs_review with the reason (for example "PPL refers to PPL004, which is no longer available"). They are never deleted or changed. The flag clears by itself when the value comes back. master_records_to_review is the admin worklist.'),
])

doc.add_heading('8. The tables', level=1)
picture('d4.png')
doc.add_heading('Changed tables', level=2)
table(['Table', 'New columns', 'Purpose'], [
    ['master_definition', 'source_system, sync_mode, control_mode, allow_st_records, lov_type_field, lov_code_field', 'Where the master comes from, how it is fetched, whether Business may control it'],
    ['master_field', 'source_system, crm_field, ref_master, lov_type', 'Which side owns the field, its name in the CRM feed, what it points to'],
    ['master_record', 'crm_key, crm_data, crm_hash, crm_synced_at, crm_status, needs_review, review_reason', 'The CRM part, change detection, removed-in-CRM state, review flag'],
    ['master_record_history', 'old_crm_data, new_crm_data', 'History keeps both parts'],
], widths=[1.6, 2.9, 2.5])
doc.add_heading('New tables', level=2)
table(['Table', 'Purpose'], [
    ['lov_type', 'Registry of lists of values: source (CRM / ST) and whether ST may extend a CRM list'],
    ['control_scope', 'The ST functions (or modules, dealers) Business can control for; data, not code'],
    ['master_control_policy', 'Control switched on for a master and a scope: deny-list or allow-list, optional approval'],
    ['master_control', "Business's choice per record and scope: state, settings, reason, who and when"],
    ['master_control_request', 'A requested change waiting for approval by a second person'],
    ['master_control_history', 'Append-only history of every control change'],
    ['master_sync_config', 'How a master is read from CRM: Solar or API, source object, key field, field mapping, schedule'],
    ['master_sync_run, master_sync_stage, master_sync_item', 'One CRM refresh: the run and its counts, the staged CRM rows, the result for each record'],
], widths=[2.4, 4.6])
doc.add_heading('Functions and views', level=2)
table(['Name', 'What it does'], [
    ['start_master_sync / apply_master_sync / fail_master_sync', 'Open a sync run, apply the staged CRM rows in one transaction, record a failed run'],
    ['set_master_control / decide_master_control_request', 'Business sets a control (or requests it when approval is needed); the approver decides'],
    ['flag_broken_references', 'Flags records that point at something that no longer exists; clears the flag when it is back'],
    ['master_available', 'What an ST function offers: ENABLED records, CRM part + ST part merged'],
    ['master_scope_state, master_new_for_review, master_record_effective, master_records_to_review', 'State of every record per scope, the Business to-do list, the merged record, the admin review list'],
], widths=[3.0, 4.0])

doc.add_heading('9. A worked example (PL / PPL, test data)', level=1)
para('This is the exact story the automated test follows, using the example from the mail: CRM holds many models and Business makes only some of them available in one ST function.')
table(['Step', 'What happens', 'Result'], [
    ['1', 'CRM sends 8 models (Altroz, Curvv EV, Curvv, Harrier, Nexon, Nexon EV, Punch, Safari).', '8 records created. Job card creation (deny-list) offers all 8. eQC Guided Check (allow-list) offers none; all 8 wait as "new, not yet enabled".'],
    ['2', 'Business enables Altroz, Harrier and Nexon for eQC Guided Check.', 'eQC offers exactly those 3; 5 still wait. Three history rows say who and why; the master version moved so dealer apps re-sync.'],
    ['3', 'A user tries to edit a CRM field, add a CRM record, or add an ST-only record to the CRM master.', 'All refused by the database. Editing the ST remark of a record is accepted.'],
    ['4', 'CRM renames Curvv EV, adds Sierra EV and drops Harrier (which Business had enabled).', '1 new, 1 changed, 6 unchanged, 1 removed. eQC now offers Altroz and Nexon. Harrier is hidden but its "enabled" choice is kept. Sierra EV waits for Business. The ST remark on Curvv EV is still there. A guided-check record that pointed at Harrier is flagged "needs review".'],
    ['5', 'CRM sends exactly the same data again.', 'Nothing is written: no new version, no history.'],
    ['6', 'CRM brings Harrier back.', 'Harrier is offered again in eQC exactly as Business had chosen; the review flag clears.'],
    ['7', 'One CRM row is malformed.', 'The other rows are applied; the run ends PARTIAL and lists the bad row.'],
    ['8', 'Business later asks to control another master (Outlets), and an approval is required for one scope.', 'One setting and one policy row; no new table. The change waits for a second person; the requester cannot approve; a rejected request changes nothing.'],
], widths=[0.5, 3.0, 3.5])

doc.add_heading('10. Decisions needed from the architect', level=1)
para('Everything below is already built with the default shown. Each is a setting or a small change, not a redesign.')
table(['#', 'Question', 'Default in this design'], [
    ['1', 'How often does CRM sync, and can CRM send only changes (a watermark)?', 'FULL and DELTA are both supported; the schedule is a column per master. Suggested: nightly FULL, plus DELTA if CRM supports it.'],
    ['2', 'What happens when CRM removes a record that Business had enabled?', 'It is hidden, kept with its history and its business choice, and restored automatically if CRM brings it back. Dependent ST records are flagged for review.'],
    ['3', 'Granularity of control: per ST function, per module or per dealer?', 'Per ST function. Module, dealer and global scopes are supported by scope_type.'],
    ['4', 'Which masters need approval of a control change, and who approves?', 'Off by default. On per master and scope, with an approver role; the requester can never approve.'],
    ['5', 'May ST add its own values to a CRM list of values?', 'No. Allowed per list (allow_st_extension) and always flagged as ST values.'],
    ['6', 'May ST users add records of their own to a CRM master?', 'No (allow_st_records = false), unless the master says so.'],
    ['7', 'Which CRM field is the stable key of each master?', 'One key_field per master in master_sync_config; record_id = that key.'],
    ['8', 'Solar or API for each master, and who runs the integration job?', 'Both supported as sync_mode and method. The job itself is not part of the database design.'],
    ['9', 'Record volumes per master and expected size of a sync.', 'Each changed record raises the master version once (same as today). For syncs of tens of thousands of rows we would batch the version bump; to be sized once volumes are known.'],
    ['10', 'Who may run a sync, set a control or approve one (roles and permissions)?', 'Not defined yet. To be mapped to app_role / role_permission.'],
    ['11', 'Dealer-owned masters: may a Dealer Admin control records for their own dealer?', 'Not defined yet. Row-level security exists for master_record; the control tables need the same rule once decided.'],
], widths=[0.4, 3.2, 3.4])

doc.add_heading('11. What the database guarantees (tested)', level=1)
bullets([
    'The CRM part of a record changes only through the CRM sync; any other writer is refused.',
    'A CRM refresh never changes a business choice or an ST field; a business choice never changes CRM data.',
    'An identical refresh writes nothing (no version bump, no history). A removed record is never deleted and comes back as it was.',
    'ST users cannot edit CRM fields, add values to a CRM list, or add records to a CRM master unless the master allows it.',
    'Control exists only for masters where it is switched on; a policy is needed for each scope; the requester can never approve; a decided request cannot be decided again.',
    'Control history is append-only; every control change raises the master version.',
    'A record that points at a removed CRM value is flagged for review, never deleted; the flag clears when the value returns.',
    'One bad CRM row does not stop the others; a DELTA run never marks anything removed.',
])
para('Proof: schema_masters_crm_test.sql (45 checks and 19 refused wrong actions) runs on PostgreSQL 16 with the existing schema tests on every change in the project.', italic=True)

doc.add_heading('12. Not covered yet', level=1)
bullets([
    'The integration job that reads Solar or the CRM API and fills master_sync_stage (needs the architect and the CRM team).',
    'The API endpoints and portal screens for control and for the review worklists (the portal today keeps data in each user\'s browser; a shared backend replaces that).',
    'Roles and permissions for sync and control (decision 10), and dealer rules for control tables (decision 11).',
    'The final list of masters and their classification: the BA teams are sending the final masters. Until then every classification is a proposal; the design does not depend on it.',
])
para('Files: docs/backend/schema_masters_crm.sql (the design), schema_masters_crm_test.sql (the tests), and the strategy note MASTER_DATA_STRATEGY.md. Test data only: no real Tata Motors data is used.', italic=True, size=9.5)

os.makedirs(OUT, exist_ok=True)
target = os.path.join(OUT, 'TML_Masters_Backend_Design.docx')
doc.save(target)
print('written', target)
