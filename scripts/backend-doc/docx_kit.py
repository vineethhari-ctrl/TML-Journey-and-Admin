"""Shared Word helpers for the backend design documents (build_doc.py = part 2 masters, build_doc_platform.py = part 3).
Call start(footer) once, then use para / bullets / table / picture on the returned document. Needs: pip install python-docx."""
import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(os.path.join(HERE, '../../build/backend-doc'))
NAVY = RGBColor(0x00, 0x22, 0x44)

doc = None


def start(footer):
    """New document with the house styles and a footer '<footer> <page number>'."""
    global doc
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(10.5)
    for name, size in (('Heading 1', 17), ('Heading 2', 13)):
        st = doc.styles[name]
        st.font.name = 'Calibri'
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = NAVY

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run(footer)
    fr.font.size = Pt(8)
    for tag, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        run = fp.add_run()
        run.font.size = Pt(8)
        if tag:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), tag)
        else:
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = txt
        run._r.append(el)
    return doc


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)


def para(text='', bold=False, italic=False, size=None, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def code(text, size=8.5):
    """A block of code (JSON, SQL) in a monospaced font."""
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Consolas'
    r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(6)
    return p


def bullets(items):
    for t in items:
        p = doc.add_paragraph(style='List Bullet')
        if isinstance(t, tuple):
            r = p.add_run(t[0])
            r.bold = True
            p.add_run(t[1])
        else:
            p.add_run(t)
        p.paragraph_format.space_after = Pt(2)


def table(head, rows, widths=None, font=9.5):
    t = doc.add_table(rows=1, cols=len(head))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font)
        r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, '002244')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(font)
            if i == 0:
                r.bold = True
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def picture(name, width=6.6):
    doc.add_picture(os.path.join(OUT, name), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def save(filename):
    os.makedirs(OUT, exist_ok=True)
    target = os.path.join(OUT, filename)
    doc.save(target)
    print('written', target)
